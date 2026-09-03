"""Relatório de Rotas por Zona de Entrega: período, filtro de zonas e exportação XLSX."""

from __future__ import annotations

from collections import defaultdict
from io import BytesIO

from django.db.models import Max

from pages.pedidos.models import Devolucao, Incidencia, TentativaEntrega, estado_segue_para_entrega
from pages.pedidos.services.relatorio_fechamento import validar_periodo
from pages.pedidos.services.zona_entrega_pedido import (
    carregar_regras_zona_por_filial,
    resolver_zona_entrega,
)
from pages.zona_entrega.models import ZonaEntrega
from sac_base.coercion import parse_date, parse_int

PERIODO_MAXIMO_ROTAS_ZONA_DIAS = 365
SEM_ZONA_LABEL = "Sem zona"
SEM_ZONA_FILTRO = "sem"

_CAMPOS_TENTATIVA = (
    "id",
    "data_tentativa",
    "carro",
    "periodo",
    "estado",
    "pedido_id",
    "motorista_id",
    "pedido__id",
    "pedido__pedido",
    "pedido__id_vonzu",
    "pedido__tipo",
    "pedido__nome_dest",
    "pedido__fone_dest",
    "pedido__fone_dest2",
    "pedido__endereco_dest",
    "pedido__cidade_dest",
    "pedido__codpost_dest",
    "pedido__volume",
    "pedido__volume_conf",
    "pedido__peso",
    "pedido__obs_rota",
    "motorista__id",
    "motorista__nome",
)


def listar_zonas_choices_filial(filial_ativa) -> list[dict]:
    if not filial_ativa:
        return []
    return [
        {"value": z.id, "label": f"{z.codigo} — {z.descricao}"}
        for z in ZonaEntrega.objects.filter(
            filial=filial_ativa,
            is_deleted=False,
            ativa=True,
        ).order_by("-prioridade", "descricao")
    ]


def _parse_zona_filtro(zonas_raw) -> tuple[set[int], bool]:
    """IDs de zona do catálogo e se o filtro inclui o grupo Sem zona (`sem`)."""
    if not isinstance(zonas_raw, list):
        zonas_raw = []
    zona_ids: set[int] = set()
    incluir_sem = False
    for x in zonas_raw:
        s = str(x).strip()
        if s.lower() == SEM_ZONA_FILTRO:
            incluir_sem = True
            continue
        v = parse_int(x)
        if v is not None:
            zona_ids.add(v)
    return zona_ids, incluir_sem


def _linha_passa_filtro_zona(zona_id, zona_ids: set[int], incluir_sem: bool) -> bool:
    if not zona_ids and not incluir_sem:
        return True
    if zona_id is None:
        return incluir_sem
    return zona_id in zona_ids


def _parse_filtros_rotas_zona(filtros: dict) -> tuple:
    """Retorna (dt_ini, dt_fim, zona_ids, incluir_sem, erro)."""
    filtros = filtros or {}
    data_inicial = str(filtros.get("data_inicial") or "").strip()
    data_final = str(filtros.get("data_final") or "").strip()
    zona_ids, incluir_sem = _parse_zona_filtro(filtros.get("zonas") or [])

    if not data_inicial:
        return None, None, zona_ids, incluir_sem, "A data inicial é obrigatória."
    if not data_final:
        return None, None, zona_ids, incluir_sem, "A data final é obrigatória."

    dt_ini = parse_date(data_inicial)
    dt_fim = parse_date(data_final)
    if dt_ini is None:
        return None, None, zona_ids, incluir_sem, "Data inicial inválida."
    if dt_fim is None:
        return None, None, zona_ids, incluir_sem, "Data final inválida."

    err_periodo = validar_periodo(dt_ini, dt_fim, max_dias=PERIODO_MAXIMO_ROTAS_ZONA_DIAS)
    if err_periodo:
        return None, None, zona_ids, incluir_sem, err_periodo
    return dt_ini, dt_fim, zona_ids, incluir_sem, None


def _formatar_periodo(dt_ini, dt_fim) -> str:
    ini = dt_ini.strftime("%d/%m/%Y")
    fim = dt_fim.strftime("%d/%m/%Y")
    if dt_ini == dt_fim:
        return ini
    return f"{ini} a {fim}"


def _peso_str(peso) -> str:
    if peso is None:
        return ""
    try:
        return str(int(peso))
    except (TypeError, ValueError):
        return str(peso)


def _nome_motorista(mov) -> str:
    if not mov.motorista_id:
        return ""
    mot = getattr(mov, "motorista", None)
    return mot.nome if mot else ""


def _peso_num(valor) -> float:
    texto = str(valor or "").strip().replace(",", ".")
    if not texto:
        return 0.0
    try:
        return float(texto)
    except (TypeError, ValueError):
        return 0.0


def _volume_pedido_num(valor_volumes) -> int:
    texto = str(valor_volumes or "").strip()
    if not texto:
        return 0
    partes = texto.split("/")
    volume_pedido = partes[1] if len(partes) >= 2 else partes[0]
    try:
        return int(str(volume_pedido).strip())
    except (TypeError, ValueError):
        return 0


def _carros_distintos(linhas: list) -> tuple[list[str], int]:
    """Números de carro distintos (ignora vazio / —), ordenados numericamente."""
    chaves: list[tuple] = []
    vistos: set = set()
    for linha in linhas:
        bruto = str(linha.get("carro") or "").strip()
        if not bruto or bruto == "—":
            continue
        try:
            chave = int(bruto)
            ordenacao = (0, chave)
        except (TypeError, ValueError):
            chave = bruto
            ordenacao = (1, bruto)
        if chave in vistos:
            continue
        vistos.add(chave)
        chaves.append((ordenacao, str(chave) if not isinstance(chave, str) else chave))
    chaves.sort(key=lambda t: t[0])
    carros = [c for _ord, c in chaves]
    return carros, len(carros)


def _agrupar_linhas_por_data(linhas: list) -> list[dict]:
    por_iso: dict[str, list] = defaultdict(list)
    for linha in linhas:
        por_iso[linha.get("data_iso") or ""].append(linha)

    dias = []
    for data_iso in sorted(por_iso.keys()):
        itens = por_iso[data_iso]
        carros, total_carros = _carros_distintos(itens)
        data_fmt = itens[0].get("data_tentativa") or data_iso
        dias.append({
            "data": data_fmt,
            "data_iso": data_iso,
            "total": len(itens),
            "peso": round(sum(_peso_num(ln.get("peso")) for ln in itens), 3),
            "volumes": sum(_volume_pedido_num(ln.get("volumes")) for ln in itens),
            "carros": carros,
            "total_carros": total_carros,
            "linhas": itens,
        })
    return dias


def montar_relatorio_rotas_zona(filial_ativa, filtros: dict) -> tuple[dict | None, str | None]:
    if not filial_ativa:
        return None, "Filial ativa não encontrada na sessão."

    dt_ini, dt_fim, zona_ids, incluir_sem, err = _parse_filtros_rotas_zona(filtros)
    if err:
        return None, err

    qs = (
        TentativaEntrega.objects
        .select_related("pedido", "motorista")
        .filter(
            data_tentativa__range=(dt_ini, dt_fim),
            pedido__filial=filial_ativa,
        )
        .only(*_CAMPOS_TENTATIVA)
        .order_by("data_tentativa", "pedido__codpost_dest", "pedido__pedido", "id")
    )
    movs = list(qs)
    regras_zona = carregar_regras_zona_por_filial(filial_ativa)
    prioridade_por_zona = {r["zona"].id: r["zona"].prioridade for r in regras_zona}

    candidatos = []
    for mov in movs:
        zona_id, zona_desc, faixa_entrega = resolver_zona_entrega(mov.pedido.codpost_dest, regras_zona)
        if not _linha_passa_filtro_zona(zona_id, zona_ids, incluir_sem):
            continue
        candidatos.append((mov, zona_id, zona_desc, faixa_entrega))

    pedido_ids = {mov.pedido_id for mov, _zid, _zd, _faixa in candidatos}

    max_dt_por_pedido = {}
    if pedido_ids:
        max_dt_por_pedido = {
            row["pedido_id"]: row["max_dt"]
            for row in (
                TentativaEntrega.objects
                .filter(pedido_id__in=pedido_ids, pedido__filial=filial_ativa)
                .values("pedido_id")
                .annotate(max_dt=Max("data_tentativa"))
            )
        }

    pedidos_com_devolucao = set()
    if pedido_ids:
        pedidos_com_devolucao = set(
            Devolucao.objects
            .filter(pedido_id__in=pedido_ids, pedido__filial=filial_ativa)
            .values_list("pedido_id", flat=True)
            .distinct()
        )

    pedidos_com_incidencia_peso_pendente = set()
    if pedido_ids:
        pedidos_com_incidencia_peso_pendente = set(
            Incidencia.objects
            .filter(
                pedido_id__in=pedido_ids,
                pedido__filial=filial_ativa,
                tipo="Peso/Volume",
                resolvido=False,
            )
            .values_list("pedido_id", flat=True)
            .distinct()
        )

    buckets: dict[int | None, list] = defaultdict(list)
    meta: dict[int | None, tuple[str, int]] = {None: (SEM_ZONA_LABEL, 0)}

    for mov, zona_id, zona_desc, faixa_entrega in candidatos:
        p = mov.pedido
        if zona_id is not None:
            meta[zona_id] = (zona_desc or SEM_ZONA_LABEL, prioridade_por_zona.get(zona_id, 0))
        tipo_abrev = "R" if (p.tipo or "").upper() == "RECOLHA" else "E"
        fones = " / ".join(f for f in [p.fone_dest or "", p.fone_dest2 or ""] if f)
        segue_para_entrega = estado_segue_para_entrega(mov.estado)
        max_dt = max_dt_por_pedido.get(p.id)
        tem_tentativa_posterior = bool(max_dt and max_dt > mov.data_tentativa)
        buckets[zona_id].append({
            "pedido_id": p.id,
            "pedido": p.pedido or str(p.id_vonzu),
            "id_vonzu": p.id_vonzu,
            "tipo": tipo_abrev,
            "nome_dest": p.nome_dest or "",
            "fones": fones,
            "endereco_dest": p.endereco_dest or "",
            "cidade_dest": p.cidade_dest or "",
            "codpost_dest": p.codpost_dest or "",
            "volumes": f"{p.volume_conf or 0}/{p.volume or 0}",
            "peso": _peso_str(p.peso),
            "periodo": mov.periodo or "",
            "data_tentativa": mov.data_tentativa.strftime("%d/%m/%Y"),
            "data_iso": mov.data_tentativa.isoformat(),
            "carro": str(mov.carro) if mov.carro is not None else "—",
            "motorista_nome": _nome_motorista(mov),
            "zona_entrega": zona_desc,
            "faixa_entrega": faixa_entrega,
            "obs_rota": p.obs_rota or "",
            "segue_para_entrega": segue_para_entrega,
            "nao_segue_para_entrega": (not segue_para_entrega) or tem_tentativa_posterior,
            "tem_devolucao": p.id in pedidos_com_devolucao,
            "tem_incidencia_peso_pendente": p.id in pedidos_com_incidencia_peso_pendente,
        })

    grupos_ord = []
    for zona_id, linhas in buckets.items():
        descricao, prioridade = meta.get(zona_id, (SEM_ZONA_LABEL, 0))
        grupos_ord.append((zona_id is None, -prioridade, descricao, zona_id, linhas))
    grupos_ord.sort(key=lambda t: (t[0], t[1], t[2]))

    grupos = []
    for _sem, _prio_neg, descricao, zona_id, linhas in grupos_ord:
        dias = _agrupar_linhas_por_data(linhas)
        grupos.append({
            "zona_id": zona_id,
            "zona": descricao,
            "total": len(linhas),
            "dias": dias,
            "linhas": linhas,
        })

    periodo_texto = _formatar_periodo(dt_ini, dt_fim)
    return {
        "grupos": grupos,
        "periodo_texto": periodo_texto,
        "data_ini_iso": dt_ini.isoformat(),
        "data_fim_iso": dt_fim.isoformat(),
    }, None


def _texto_resumo_dia(zona: str, dia: dict) -> str:
    carros = dia.get("carros") or []
    carros_txt = ", ".join(carros) if carros else "—"
    return (
        f"{dia.get('data') or ''} • {zona}: {dia.get('total', 0)} pedido(s) • "
        f"{dia.get('peso', 0)} kg • {dia.get('volumes', 0)} vol • "
        f"Carros {carros_txt} • Total de carros {dia.get('total_carros', 0)}"
    )


def _ajustar_largura_colunas(ws) -> None:
    for col in ws.columns:
        max_len = 0
        letter = col[0].column_letter
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 48)


def gerar_xlsx_relatorio_rotas_zona(payload: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    font_bold = Font(bold=True)
    fill_nao_segue = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")
    fill_resumo = PatternFill(start_color="D6E3F0", end_color="D6E3F0", fill_type="solid")
    periodo_texto = payload.get("periodo_texto") or ""
    grupos = payload.get("grupos") or []

    ws_resumo = wb.active
    ws_resumo.title = "Resumo"
    if periodo_texto:
        ws_resumo.append([f"Rotas por Zona — resumo por data — {periodo_texto}"])
        ws_resumo["A1"].font = font_bold
        ws_resumo.append([])
    ws_resumo.append([
        "Zona",
        "Data",
        "Pedidos",
        "Peso",
        "Volumes",
        "Carros",
        "Total de carros",
    ])
    for cell in ws_resumo[ws_resumo.max_row]:
        cell.font = font_bold

    for grupo in grupos:
        zona_label = (grupo.get("zona") or "").strip() or SEM_ZONA_LABEL
        dias = grupo.get("dias") or _agrupar_linhas_por_data(grupo.get("linhas") or [])
        for dia in dias:
            carros = dia.get("carros") or []
            ws_resumo.append([
                zona_label,
                dia.get("data") or "",
                dia.get("total", 0),
                dia.get("peso", 0),
                dia.get("volumes", 0),
                ", ".join(carros) if carros else "—",
                dia.get("total_carros", 0),
            ])
    _ajustar_largura_colunas(ws_resumo)

    ws = wb.create_sheet("Detalhe")
    if periodo_texto:
        ws.append([f"Rotas por Zona — {periodo_texto}"])
        ws["A1"].font = font_bold
        ws.append([])

    cabecalhos = [
        "Zona",
        "Data",
        "Carro",
        "Motorista",
        "Referência",
        "T",
        "Destinatário",
        "Telefone(s)",
        "Endereço",
        "Cidade",
        "C. Postal",
        "Vol",
        "Peso",
        "Per.",
        "Obs. Rota",
        "Faixa",
        "Devolução",
        "Segue para entrega",
    ]
    ws.append(cabecalhos)
    for cell in ws[ws.max_row]:
        cell.font = font_bold

    for grupo in grupos:
        zona_label = (grupo.get("zona") or "").strip() or SEM_ZONA_LABEL
        dias = grupo.get("dias") or _agrupar_linhas_por_data(grupo.get("linhas") or [])
        for dia in dias:
            resumo_row = ws.max_row + 1
            ws.append([_texto_resumo_dia(zona_label, dia)])
            for cell in ws[resumo_row]:
                cell.font = font_bold
                cell.fill = fill_resumo
            for linha in dia.get("linhas") or []:
                ref = linha.get("pedido") or ""
                if linha.get("tem_devolucao"):
                    ref = f"{ref} (Dev)"
                row_idx = ws.max_row + 1
                ws.append([
                    zona_label,
                    linha.get("data_tentativa") or "",
                    linha.get("carro") or "",
                    linha.get("motorista_nome") or "",
                    ref,
                    linha.get("tipo") or "",
                    linha.get("nome_dest") or "",
                    linha.get("fones") or "",
                    linha.get("endereco_dest") or "",
                    linha.get("cidade_dest") or "",
                    linha.get("codpost_dest") or "",
                    linha.get("volumes") or "",
                    linha.get("peso") or "",
                    linha.get("periodo") or "",
                    linha.get("obs_rota") or "",
                    linha.get("faixa_entrega") or "",
                    "Sim" if linha.get("tem_devolucao") else "Não",
                    "Não" if linha.get("nao_segue_para_entrega") else "Sim",
                ])
                if linha.get("nao_segue_para_entrega"):
                    for cell in ws[row_idx]:
                        cell.fill = fill_nao_segue

    _ajustar_largura_colunas(ws)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
