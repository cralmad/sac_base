"""Relatório de Rotas por Carro/Motorista: montagem de grupos e exportação XLSX."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO
from itertools import groupby

from pages.pedidos.models import Devolucao, Incidencia, TentativaEntrega, tentativa_segue_para_entrega
from pages.pedidos.services.zona_entrega_pedido import (
    carregar_regras_zona_por_filial,
    resolver_zona_e_faixa_entrega,
)
from sac_base.smart_filter import apply_smart_number_filter


def _parse_filtros_rotas(filtros: dict) -> tuple[date | None, str, list[int], str, str | None]:
    """Extrai e valida filtros. Retorna (dt, carro, motorista_ids, agrupamento, erro)."""
    filtros = filtros or {}
    data_tentativa = str(filtros.get("data_tentativa") or "").strip()
    carro_filtro = str(filtros.get("carro") or "").strip()
    motorista_ids_raw = filtros.get("motoristas") or []
    if not isinstance(motorista_ids_raw, list):
        motorista_ids_raw = []
    motorista_ids = [int(v) for v in motorista_ids_raw if str(v).isdigit()]
    agrupamento = str(filtros.get("agrupamento") or "carro").strip().lower()
    if agrupamento not in ("carro", "motorista"):
        agrupamento = "carro"

    if not data_tentativa:
        return None, carro_filtro, motorista_ids, agrupamento, "A data é obrigatória."

    try:
        dt = datetime.strptime(data_tentativa, "%Y-%m-%d").date()
    except ValueError:
        return None, carro_filtro, motorista_ids, agrupamento, "Data inválida."

    return dt, carro_filtro, motorista_ids, agrupamento, None


def montar_relatorio_rotas(filial_ativa, filtros: dict) -> tuple[dict | None, str | None]:
    """Monta grupos de rotas para a filial. Retorna (payload, None) ou (None, mensagem_erro)."""
    if not filial_ativa:
        return None, "Filial ativa não encontrada na sessão."

    dt, carro_filtro, motorista_ids, agrupamento, err = _parse_filtros_rotas(filtros)
    if err:
        return None, err

    qs = (
        TentativaEntrega.objects
        .select_related("pedido", "motorista")
        .filter(data_tentativa=dt, pedido__filial=filial_ativa)
    )
    if carro_filtro:
        qs = apply_smart_number_filter(qs, "carro", carro_filtro)
    if motorista_ids:
        qs = qs.filter(motorista_id__in=motorista_ids)

    if agrupamento == "motorista":
        qs = qs.order_by("motorista__nome", "pedido__codpost_dest", "pedido__pedido")
        key_fn = lambda m: (m.motorista_id, m.motorista.nome if m.motorista_id else "")
    else:
        qs = qs.order_by("carro", "pedido__codpost_dest", "pedido__pedido")
        key_fn = lambda m: m.carro

    movs = list(qs)
    pedido_ids = {m.pedido_id for m in movs}
    pedidos_com_tentativa_posterior = set()
    if pedido_ids:
        pedidos_com_tentativa_posterior = set(
            TentativaEntrega.objects
            .filter(
                pedido_id__in=pedido_ids,
                data_tentativa__gt=dt,
                pedido__filial=filial_ativa,
            )
            .values_list("pedido_id", flat=True)
            .distinct()
        )

    pedidos_com_devolucao = set()
    if pedido_ids:
        pedidos_com_devolucao = set(
            Devolucao.objects
            .filter(pedido_id__in=pedido_ids, pedido__filial=filial_ativa)
            .values_list("pedido_id", flat=True)
            .distinct()
        )

    pedidos_com_incidencia_peso_pendente = set()
    if pedido_ids:
        pedidos_com_incidencia_peso_pendente = set(
            Incidencia.objects
            .filter(
                pedido_id__in=pedido_ids,
                pedido__filial=filial_ativa,
                tipo="Peso/Volume",
                resolvido=False,
            )
            .values_list("pedido_id", flat=True)
            .distinct()
        )

    regras_zona = carregar_regras_zona_por_filial(filial_ativa)

    grupos = []
    for grupo_key, items in groupby(movs, key=key_fn):
        linhas = []
        data_str = None
        motorista_nome = ""
        carro_val = None
        for mov in items:
            p = mov.pedido
            if data_str is None:
                data_str = mov.data_tentativa.strftime("%d/%m/%Y")
            if agrupamento == "motorista":
                motorista_nome = mov.motorista.nome if mov.motorista_id else ""
                carro_val = None
            else:
                carro_val = mov.carro
            tipo_abrev = "R" if (p.tipo or "").upper() == "RECOLHA" else "E"
            fones = " / ".join(f for f in [p.fone_dest or "", p.fone_dest2 or ""] if f)
            peso_str = ""
            if p.peso is not None:
                try:
                    peso_str = str(int(p.peso))
                except (TypeError, ValueError):
                    peso_str = str(p.peso)
            segue_para_entrega = tentativa_segue_para_entrega(mov)
            tem_tentativa_posterior = p.id in pedidos_com_tentativa_posterior
            zona_entrega, faixa_entrega = resolver_zona_e_faixa_entrega(p.codpost_dest, regras_zona)
            linhas.append({
                "pedido_id": p.id,
                "pedido": p.pedido or str(p.id_vonzu),
                "id_vonzu": p.id_vonzu,
                "tipo": tipo_abrev,
                "nome_dest": p.nome_dest or "",
                "fones": fones,
                "endereco_dest": p.endereco_dest or "",
                "cidade_dest": p.cidade_dest or "",
                "codpost_dest": p.codpost_dest or "",
                "volumes": f"{p.volume_conf or 0}/{p.volume or 0}",
                "peso": peso_str,
                "periodo": mov.periodo or "",
                "zona_entrega": zona_entrega,
                "faixa_entrega": faixa_entrega,
                "obs_rota": p.obs_rota or "",
                "segue_para_entrega": segue_para_entrega,
                "nao_segue_para_entrega": (not segue_para_entrega) or tem_tentativa_posterior,
                "tem_devolucao": p.id in pedidos_com_devolucao,
                "tem_incidencia_peso_pendente": p.id in pedidos_com_incidencia_peso_pendente,
            })
        grupos.append({
            "carro": str(carro_val) if carro_val is not None else "—",
            "motorista_nome": motorista_nome,
            "data_tentativa": data_str or dt.strftime("%d/%m/%Y"),
            "total": len(linhas),
            "linhas": linhas,
        })

    return {
        "grupos": grupos,
        "data_fmt": dt.strftime("%d/%m/%Y"),
        "agrupamento": agrupamento,
        "data_iso": dt.isoformat(),
    }, None


def gerar_xlsx_relatorio_rotas(payload: dict) -> bytes:
    """Gera workbook .xlsx com as linhas do relatório de rotas (colunas alinhadas à UI)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    agrupamento = payload.get("agrupamento") or "carro"
    ws.title = "Rotas por Motorista" if agrupamento == "motorista" else "Rotas por Carro"

    data_fmt = payload.get("data_fmt") or ""
    if data_fmt:
        titulo = (
            f"Rotas por Motorista — {data_fmt}"
            if agrupamento == "motorista"
            else f"Rotas por Carro — {data_fmt}"
        )
        ws.append([titulo])
        ws["A1"].font = Font(bold=True)
        ws.append([])

    cabecalhos = [
        "Carro" if agrupamento != "motorista" else "Motorista",
        "Data",
        "Referência",
        "T",
        "Destinatário",
        "Telefone(s)",
        "Endereço",
        "Cidade",
        "C. Postal",
        "Vol",
        "Peso",
        "Per.",
        "Obs. Rota",
        "Zona",
        "Faixa",
        "Devolução",
        "Segue para entrega",
    ]
    ws.append(cabecalhos)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True)

    fill_nao_segue = PatternFill(start_color="E0E0E0", end_color="E0E0E0", fill_type="solid")

    for grupo in payload.get("grupos") or []:
        if agrupamento == "motorista":
            grupo_label = (grupo.get("motorista_nome") or "").strip() or "Sem motorista"
        else:
            carro = grupo.get("carro")
            grupo_label = f"Carro {carro}" if carro not in (None, "", "—") else "Sem carro"
        data_grupo = grupo.get("data_tentativa") or data_fmt

        for linha in grupo.get("linhas") or []:
            ref = linha.get("pedido") or ""
            if linha.get("tem_devolucao"):
                ref = f"{ref} (Dev)"
            row_idx = ws.max_row + 1
            ws.append([
                grupo_label,
                data_grupo,
                ref,
                linha.get("tipo") or "",
                linha.get("nome_dest") or "",
                linha.get("fones") or "",
                linha.get("endereco_dest") or "",
                linha.get("cidade_dest") or "",
                linha.get("codpost_dest") or "",
                linha.get("volumes") or "",
                linha.get("peso") or "",
                linha.get("periodo") or "",
                linha.get("obs_rota") or "",
                linha.get("zona_entrega") or "",
                linha.get("faixa_entrega") or "",
                "Sim" if linha.get("tem_devolucao") else "Não",
                "Não" if linha.get("nao_segue_para_entrega") else "Sim",
            ])
            if linha.get("nao_segue_para_entrega"):
                for cell in ws[row_idx]:
                    cell.fill = fill_nao_segue

    for col in ws.columns:
        max_len = 0
        letter = col[0].column_letter
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 48)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
