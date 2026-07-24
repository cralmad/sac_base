"""Importação de pedidos a partir de XLSX ENOVO."""

from __future__ import annotations

import io
import zipfile
from datetime import date, datetime
from decimal import Decimal
from xml.etree import ElementTree as ET

from django.utils import timezone

from pages.cad_cliente.models import Cliente
from pages.motorista.models import Motorista
from sac_base.coercion import parse_date, parse_datetime, parse_decimal, parse_int

from .importador_persistencia import (
    coletar_pedidos_movimentacao_ausentes_no_arquivo,
    gerar_relatorio_importacao,
    montar_dados_volumes_agrupados,
    persistir_ou_erro,
    validar_data_unica_para_analise,
)
from .normalizacao import normalizar_estado_enovo

_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

COD_SERVICO_TIPO = {
    "DIST": "ENTREGA",
}

LAYOUT_COMPLETO = "completo"

HEADERS_OBRIGATORIOS = (
    "Tipo",
    "TRK",
    "Cod. Serviço",
    "Referência",
    "Cod. Cliente",
    "Remetente",
    "Morada Remetente",
    "CP Remetente",
    "Localidade Remetente",
    "Destinatário",
    "Morada Destinatário",
    "CP Destinatário",
    "Localidade Destinatário",
    "Último Estado",
    "Data Último Estado",
    "Data Descarga",
    "Data Criação",
    "Volumes",
    "Peso (Kg)",
)

# Tokens na Referência (case-insensitive) que forçam tipo RECOLHA
_TOKENS_RECOLHA = ("cd", "lmpt")


def tipo_por_referencia_recolha(referencia: str | None) -> bool:
    """True se Referência contém 'cd' ou 'lmpt' (ordem irrelevante, case-insensitive)."""
    if not referencia:
        return False
    ref = str(referencia).casefold()
    return any(token in ref for token in _TOKENS_RECOLHA)


def resolver_tipo(cod_servico: str | None, referencia: str | None) -> str | None:
    if tipo_por_referencia_recolha(referencia):
        return "RECOLHA"
    cod = (cod_servico or "").strip().upper()
    return COD_SERVICO_TIPO.get(cod)


def _cell_col_index(cell_ref: str) -> int:
    col = ""
    for ch in cell_ref:
        if ch.isalpha():
            col += ch
        else:
            break
    n = 0
    for ch in col:
        n = n * 26 + (ord(ch.upper()) - 64)
    return n - 1


def _ler_shared_strings(zf: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in zf.namelist():
        return []
    root = ET.fromstring(zf.read("xl/sharedStrings.xml"))
    strings = []
    for si in root.findall("m:si", _NS):
        texts = [t.text or "" for t in si.findall(".//m:t", _NS)]
        strings.append("".join(texts))
    return strings


def _cell_value(cell, shared_strings: list[str]):
    t = cell.get("t")
    v = cell.find("m:v", _NS)
    if v is None:
        is_el = cell.find("m:is", _NS)
        if is_el is not None:
            return "".join(x.text or "" for x in is_el.findall(".//m:t", _NS))
        return None
    raw = v.text
    if t == "s":
        return shared_strings[int(raw)]
    if t == "b":
        return raw == "1"
    if t == "inlineStr":
        return raw
    # número ou data serial — devolver string bruta; datas costumam vir como shared/string no export ENOVO
    return raw


def _sheet_path(zf: zipfile.ZipFile) -> str:
    for name in zf.namelist():
        if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
            return name
    raise ValueError("Planilha não encontrada no XLSX.")


def _parse_xlsx_via_xml(conteudo_bytes: bytes) -> list[list]:
    with zipfile.ZipFile(io.BytesIO(conteudo_bytes)) as zf:
        shared = _ler_shared_strings(zf)
        sheet = ET.fromstring(zf.read(_sheet_path(zf)))
        rows_out = []
        for row in sheet.findall("m:sheetData/m:row", _NS):
            cells = {}
            max_c = -1
            for c in row.findall("m:c", _NS):
                idx = _cell_col_index(c.get("r", "A1"))
                cells[idx] = _cell_value(c, shared)
                max_c = max(max_c, idx)
            rows_out.append([cells.get(i) for i in range(max_c + 1)])
        return rows_out


def _parse_xlsx_via_openpyxl(conteudo_bytes: bytes) -> list[list]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(conteudo_bytes), data_only=True, read_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        return [list(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _parse_xlsx_bytes(conteudo_bytes: bytes) -> list[tuple[int, dict]]:
    """Retorna lista (num_linha_1based_excel, dict_por_header)."""
    try:
        rows = _parse_xlsx_via_xml(conteudo_bytes)
    except Exception:
        try:
            rows = _parse_xlsx_via_openpyxl(conteudo_bytes)
        except Exception as exc:
            raise ValueError("Não foi possível ler o XLSX.") from exc

    if not rows:
        return []

    headers_raw = rows[0]
    headers = []
    seen = {}
    for h in headers_raw:
        name = str(h).strip() if h is not None else ""
        if name in seen:
            seen[name] += 1
            name = f"{name}__{seen[name]}"
        else:
            seen[name] = 0
        headers.append(name)

    result = []
    for i, row in enumerate(rows[1:], start=2):
        if row is None or all(c is None or str(c).strip() == "" for c in row):
            continue
        data = {}
        for idx, header in enumerate(headers):
            data[header] = row[idx] if idx < len(row) else None
        result.append((i, data))
    return result


def _as_text(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(valor, date):
        return valor.isoformat()
    return str(valor).strip()


def _parse_date_flex(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = _as_text(valor)
    if not texto:
        return None
    # "2026-07-23 00:00" → primeiros 10 chars
    if len(texto) >= 10 and texto[4] == "-" and texto[7] == "-":
        return parse_date(texto[:10])
    return parse_date(texto)


def _parse_datetime_flex(valor):
    if isinstance(valor, datetime):
        dt = valor
        return timezone.make_aware(dt) if timezone.is_naive(dt) else dt
    texto = _as_text(valor)
    return parse_datetime(texto, context="csv")


def _concat_obs(*partes) -> str | None:
    bits = []
    for p in partes:
        t = _as_text(p)
        if t:
            bits.append(t)
    return " | ".join(bits) if bits else None


def _tem_header(presentes: set, nome: str) -> bool:
    if nome in presentes:
        return True
    return any(h == nome or h.startswith(f"{nome}__") for h in presentes)


def _validar_headers(primeira_linha: dict, obrigatorios=HEADERS_OBRIGATORIOS) -> list[str]:
    presentes = set(primeira_linha.keys())
    faltando = [h for h in obrigatorios if not _tem_header(presentes, h)]
    if faltando:
        return [f"Colunas obrigatórias em falta: {', '.join(faltando)}"]
    return []


def _get_col(row: dict, nome: str):
    if nome in row:
        return row[nome]
    for k, v in row.items():
        if k.startswith(f"{nome}__"):
            continue
    return row.get(nome)


def _linha_eh_encargo(row: dict) -> bool:
    """True se a coluna A (Tipo) for ENCARGO — linha a ignorar."""
    return _as_text(_get_col(row, "Tipo")).casefold() == "encargo"


def _normalizar_linha_bruta(num_linha: int, row: dict):
    """Normaliza uma linha física do XLSX (ainda não agregada)."""
    erros = []

    trk_raw = _as_text(_get_col(row, "TRK"))
    try:
        id_vonzu = int(trk_raw)
    except (ValueError, TypeError):
        erros.append(f"Linha {num_linha}: TRK inválido — '{trk_raw}'")
        return None, erros

    referencia = _as_text(_get_col(row, "Referência")) or None
    cod_servico = _as_text(_get_col(row, "Cod. Serviço"))
    tipo = resolver_tipo(cod_servico, referencia)
    if not tipo:
        erros.append(
            f"Linha {num_linha}: Cod. Serviço inválido — '{cod_servico}' "
            "(sem mapeamento e Referência sem cd/lmpt)"
        )

    estado_raw = _as_text(_get_col(row, "Último Estado"))
    estado = normalizar_estado_enovo(estado_raw)
    if estado is None:
        erros.append(
            f"Linha {num_linha}: Último Estado não mapeado — '{estado_raw or '(vazio)'}'"
        )

    criado = _parse_datetime_flex(_get_col(row, "Data Criação"))
    if criado is None:
        erros.append(
            f"Linha {num_linha}: 'Data Criação' inválida — "
            f"'{_as_text(_get_col(row, 'Data Criação'))}'"
        )

    atualizacao = _parse_datetime_flex(_get_col(row, "Data Último Estado"))
    if atualizacao is None:
        erros.append(
            f"Linha {num_linha}: 'Data Último Estado' inválida — "
            f"'{_as_text(_get_col(row, 'Data Último Estado'))}'"
        )

    prev_entrega = _parse_date_flex(_get_col(row, "Data Descarga"))
    if prev_entrega is None:
        erros.append(
            f"Linha {num_linha}: 'Data Descarga' inválida — "
            f"'{_as_text(_get_col(row, 'Data Descarga'))}'"
        )

    cliente_raw = _as_text(_get_col(row, "Cod. Cliente"))
    cliente_pk = parse_int(cliente_raw, context="csv")
    if cliente_pk is None:
        erros.append(f"Linha {num_linha}: Cod. Cliente inválido — '{cliente_raw}'")

    qtd = parse_int(_as_text(_get_col(row, "Volumes")), context="csv")
    if qtd is None:
        qtd = 0

    peso = parse_decimal(_as_text(_get_col(row, "Peso (Kg)")), context="csv")
    if peso is None:
        peso = Decimal("0")

    if erros:
        return None, erros

    # RECOLHA (cd/lmpt): operação inversa — dados operacionais vêm do Remetente.
    if tipo == "RECOLHA":
        nome_dest = _as_text(_get_col(row, "Remetente")) or None
        email_dest = None
        fone_dest = _as_text(_get_col(row, "Contacto Remetente")) or None
        endereco_dest = _as_text(_get_col(row, "Morada Remetente")) or None
        codpost_dest = _as_text(_get_col(row, "CP Remetente")) or None
        cidade_dest = _as_text(_get_col(row, "Localidade Remetente")) or None
    else:
        nome_dest = _as_text(_get_col(row, "Destinatário")) or None
        email_dest = _as_text(_get_col(row, "E-mail Destinatário")) or None
        fone_dest = _as_text(_get_col(row, "Contacto Destinatário")) or None
        endereco_dest = _as_text(_get_col(row, "Morada Destinatário")) or None
        codpost_dest = _as_text(_get_col(row, "CP Destinatário")) or None
        cidade_dest = _as_text(_get_col(row, "Localidade Destinatário")) or None

    return {
        "id_vonzu": id_vonzu,
        "pedido": referencia,
        "tipo": tipo,
        "criado": criado,
        "atualizacao": atualizacao,
        "prev_entrega": prev_entrega,
        "dt_entrega": _parse_date_flex(_get_col(row, "Data Entrega")),
        "estado": estado,
        "qtd": qtd,
        "peso": peso,
        "nome_dest": nome_dest,
        "email_dest": email_dest,
        "fone_dest": fone_dest,
        "fone_dest2": None,
        "endereco_dest": endereco_dest,
        "codpost_dest": codpost_dest,
        "cidade_dest": cidade_dest,
        "obs": _concat_obs(
            _get_col(row, "Observações Estado"),
            _get_col(row, "Observações Carga"),
            _get_col(row, "Observações Descarga"),
            _get_col(row, "Obs. Internas"),
        ),
        "expresso": False,
        "cliente_pk": cliente_pk,
        "id_enovo_motorista": _as_text(_get_col(row, "Cód. Motorista")) or None,
        "description_raw": _as_text(_get_col(row, "Descrição")),
        "_num_linha": num_linha,
    }, []


def _agregar_por_trk(linhas_norm: list[tuple[int, dict]]):
    """Agrega linhas físicas pelo TRK (id_vonzu). Primeira ocorrência define cabeçalho."""
    ordem = []
    por_trk = {}
    linhas_extras = 0

    for num_linha, dados in linhas_norm:
        key = dados["id_vonzu"]
        if key not in por_trk:
            agg = dict(dados)
            agg["volume"] = int(dados["qtd"] or 0)
            agg["peso"] = Decimal(dados["peso"] or 0)
            desc = (dados.get("description_raw") or "").strip()
            agg["description_raw"] = desc
            agg.pop("qtd", None)
            por_trk[key] = (num_linha, agg)
            ordem.append(key)
        else:
            linhas_extras += 1
            _nl, agg = por_trk[key]
            agg["volume"] = int(agg["volume"] or 0) + int(dados["qtd"] or 0)
            agg["peso"] = Decimal(agg["peso"] or 0) + Decimal(dados["peso"] or 0)
            desc = (dados.get("description_raw") or "").strip()
            if desc:
                if agg["description_raw"]:
                    agg["description_raw"] = f"{agg['description_raw']}; {desc}"
                else:
                    agg["description_raw"] = desc

    agregadas = [por_trk[k] for k in ordem]
    return agregadas, linhas_extras


def _resolver_fks_estrito(filial, linhas_agg):
    """Resolve cliente (PK) e motorista (id_enovo). Falhas → erros (abort)."""
    cliente_pks = {d["cliente_pk"] for _, d in linhas_agg if d.get("cliente_pk") is not None}
    ids_enovo = {
        d["id_enovo_motorista"]
        for _, d in linhas_agg
        if d.get("id_enovo_motorista")
    }

    clientes_ok = set(
        Cliente.objects.filter(pk__in=cliente_pks, is_deleted=False).values_list("pk", flat=True)
    )
    motoristas_map = {
        m.id_enovo: m.id
        for m in Motorista.objects.filter(
            filial=filial,
            id_enovo__in=ids_enovo,
            is_deleted=False,
        )
        if m.id_enovo
    }

    erros = []
    resultado = []
    for num_linha, dados in linhas_agg:
        dados = dict(dados)
        cliente_pk = dados.pop("cliente_pk", None)
        id_enovo_mot = dados.pop("id_enovo_motorista", None)

        if cliente_pk not in clientes_ok:
            erros.append(
                f"Linha {num_linha}: Cod. Cliente={cliente_pk} inexistente ou inactivo"
            )
            continue

        motorista_id = None
        if id_enovo_mot:
            motorista_id = motoristas_map.get(id_enovo_mot)
            if motorista_id is None:
                erros.append(
                    f"Linha {num_linha}: Cód. Motorista '{id_enovo_mot}' "
                    "sem Motorista.id_enovo correspondente nesta filial"
                )
                continue

        dados["cliente_id"] = cliente_pk
        dados["motorista_id"] = motorista_id
        resultado.append((num_linha, dados))

    return resultado, erros


def importar_xlsx(
    conteudo_bytes,
    filial,
    nome_arquivo,
    analisar_movimentacoes_dia=False,
    forcar_atualizacao=False,
):
    """Importa XLSX ENOVO (Listagem de Envios) para a filial."""
    try:
        linhas_raw = _parse_xlsx_bytes(conteudo_bytes)
    except Exception:
        return {
            "sucesso": False,
            "erros": ["Erro ao ler XLSX: ficheiro inválido ou corrompido."],
            "relatorio": "",
            "stats": {},
        }

    if not linhas_raw:
        return {
            "sucesso": False,
            "erros": ["O arquivo XLSX está vazio ou sem dados."],
            "relatorio": "",
            "stats": {},
        }

    erros_header = _validar_headers(linhas_raw[0][1])
    if erros_header:
        return {"sucesso": False, "erros": erros_header, "relatorio": "", "stats": {}}

    total_lidas = len(linhas_raw)
    todos_erros = []
    linhas_norm = []
    ignorados_encargo = 0
    for num_linha, row in linhas_raw:
        if _linha_eh_encargo(row):
            ignorados_encargo += 1
            continue
        dados, erros = _normalizar_linha_bruta(num_linha, row)
        if erros:
            todos_erros.extend(erros)
        else:
            linhas_norm.append((num_linha, dados))

    if todos_erros:
        return {"sucesso": False, "erros": todos_erros, "relatorio": "", "stats": {}}

    if not linhas_norm:
        return {
            "sucesso": False,
            "erros": [
                "Nenhuma linha elegível para importação "
                f"(ignoradas ENCARGO: {ignorados_encargo})."
            ],
            "relatorio": "",
            "stats": {},
        }

    linhas_agg, ignoradas = _agregar_por_trk(linhas_norm)
    linhas_resolvidas, erros_fk = _resolver_fks_estrito(filial, linhas_agg)
    if erros_fk:
        return {"sucesso": False, "erros": erros_fk, "relatorio": "", "stats": {}}

    data_analise_movimentacao = None
    pedidos_mov_ausentes_no_arquivo = []
    if analisar_movimentacoes_dia:
        data_analise_movimentacao = validar_data_unica_para_analise(linhas_resolvidas)
        if not data_analise_movimentacao:
            return {
                "sucesso": False,
                "erros": [
                    "Para analisar movimentações do dia, o arquivo deve conter "
                    "pedidos com apenas uma data válida em 'Data Descarga'.",
                ],
                "relatorio": "",
                "stats": {},
            }

    ok, resultado = persistir_ou_erro(
        filial,
        linhas_resolvidas,
        match_por_referencia=True,
        forcar_atualizacao=forcar_atualizacao,
    )
    if not ok:
        return {"sucesso": False, "erros": resultado, "relatorio": "", "stats": {}}

    criados = resultado["criados"]
    atualizados = resultado["atualizados"]
    sem_alteracao = resultado["sem_alteracao"]
    tentativas = resultado["tentativas"]
    remapeamentos = resultado.get("remapeamentos") or []

    if analisar_movimentacoes_dia:
        id_vonzus_importados = {dados["id_vonzu"] for _, dados in linhas_resolvidas}
        pedidos_mov_ausentes_no_arquivo = coletar_pedidos_movimentacao_ausentes_no_arquivo(
            filial=filial,
            data_base=data_analise_movimentacao,
            id_vonzus_importados=id_vonzus_importados,
        )

    relatorio = gerar_relatorio_importacao(
        nome_arquivo=nome_arquivo,
        filial=filial,
        total_lidas=total_lidas,
        ignoradas=ignoradas + ignorados_encargo,
        criados=criados,
        atualizados=atualizados,
        sem_alteracao=sem_alteracao,
        tentativas=tentativas,
        avisos=[],
        analise_movimentacao_ativada=analisar_movimentacoes_dia,
        data_analise=data_analise_movimentacao,
        pedidos_mov_ausentes_no_arquivo=pedidos_mov_ausentes_no_arquivo,
        remapeamentos=remapeamentos,
    )

    dados_volumes, _ = montar_dados_volumes_agrupados(linhas_resolvidas)

    return {
        "sucesso": True,
        "erros": [],
        "relatorio": relatorio,
        "stats": {
            "total_lidas": total_lidas,
            "ignoradas": ignoradas,
            "ignorados_encargo": ignorados_encargo,
            "criados": criados,
            "atualizados": atualizados,
            "sem_alteracao": sem_alteracao,
            "tentativas": tentativas,
            "avisos_fk": 0,
            "ids_remapeados": len(remapeamentos),
            "layout": LAYOUT_COMPLETO,
        },
        "dados_volumes": dados_volumes,
    }


HEADERS_ARTIGOS_ENOVO = (
    "TRK",
)


def _normalizar_linha_artigos_enovo(num_linha: int, row: dict):
    """Normalização mínima para importação de artigos (sem validar estado/FKs)."""
    erros = []
    trk_raw = _as_text(_get_col(row, "TRK"))
    try:
        id_vonzu = int(trk_raw)
    except (ValueError, TypeError):
        erros.append(f"Linha {num_linha}: TRK inválido — '{trk_raw}'")
        return None, erros

    volume = parse_int(_as_text(_get_col(row, "Volumes")), context="csv")
    peso = parse_decimal(_as_text(_get_col(row, "Peso (Kg)")), context="csv")

    return {
        "id_vonzu": id_vonzu,
        "pedido": _as_text(_get_col(row, "Referência")) or None,
        "description_raw": _as_text(_get_col(row, "Descrição")),
        "volume": volume,
        "peso": peso,
    }, erros


def parse_xlsx_artigos_sem_persistir(conteudo_bytes):
    """Parseia XLSX ENOVO e retorna artigos sem persistir.

    Coluna ``Descrição`` é opcional (export actual pode não a trazer).
    Reutiliza ``montar_dados_volumes_agrupados`` + ``parse_description``.
    Não altera o fluxo CSV.
    """
    try:
        linhas_raw = _parse_xlsx_bytes(conteudo_bytes)
    except Exception:
        return {
            "sucesso": False,
            "erros": ["Erro ao ler XLSX: ficheiro inválido ou corrompido."],
            "pedidos": [],
        }

    if not linhas_raw:
        return {
            "sucesso": False,
            "erros": ["O arquivo XLSX está vazio ou sem dados."],
            "pedidos": [],
        }

    erros_header = _validar_headers(linhas_raw[0][1], HEADERS_ARTIGOS_ENOVO)
    if erros_header:
        return {"sucesso": False, "erros": erros_header, "pedidos": []}

    todos_erros = []
    linhas_norm = []
    for num_linha, row in linhas_raw:
        if _linha_eh_encargo(row):
            continue
        dados, erros = _normalizar_linha_artigos_enovo(num_linha, row)
        if erros:
            todos_erros.extend(erros)
        else:
            linhas_norm.append((num_linha, dados))

    if todos_erros:
        return {"sucesso": False, "erros": todos_erros, "pedidos": []}

    dados_volumes, _detalhes = montar_dados_volumes_agrupados(linhas_norm)
    return {
        "sucesso": True,
        "erros": [],
        "pedidos": dados_volumes,
    }
