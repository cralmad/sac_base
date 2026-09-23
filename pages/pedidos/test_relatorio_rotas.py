"""Relatório de Rotas — montagem de grupos e exportação XLSX."""

import importlib.util
from datetime import date
from io import BytesIO

from django.test import TestCase
from django.utils import timezone

from pages.core.models import Pais
from pages.filial.models import Filial
from pages.pedidos.models import Incidencia, Pedido, TentativaEntrega
from pages.pedidos.services.relatorio_rotas import (
    gerar_xlsx_relatorio_rotas,
    montar_relatorio_rotas,
)


class RelatorioRotasServiceTests(TestCase):
    d0 = date(2026, 7, 24)

    @classmethod
    def setUpTestData(cls):
        cls.pais = Pais.objects.create(nome="PT-RR", sigla="PT", codigo_tel="+351")
        cls.filial = Filial.objects.create(
            codigo="RRCL",
            nome="FIL RR ROTAS",
            pais_atuacao=cls.pais,
            is_matriz=True,
        )

    def _pedido(self, id_vonzu: int, pedido_ref: str = "") -> Pedido:
        now = timezone.now()
        return Pedido.objects.create(
            filial=self.filial,
            id_vonzu=id_vonzu,
            pedido=pedido_ref or str(id_vonzu),
            tipo="ENTREGA",
            nome_dest="Dest Teste",
            criado=now,
            atualizacao=now,
        )

    def _tentativa(self, pedido: Pedido, carro: int = 1) -> TentativaEntrega:
        return TentativaEntrega.objects.create(
            pedido=pedido,
            data_tentativa=self.d0,
            carro=carro,
            periodo="MANHA",
            estado="created",
        )

    def test_montar_exige_filial(self):
        payload, err = montar_relatorio_rotas(None, {"data_tentativa": self.d0.isoformat()})
        self.assertIsNone(payload)
        self.assertIn("Filial ativa", err)

    def test_montar_exige_data(self):
        payload, err = montar_relatorio_rotas(self.filial, {})
        self.assertIsNone(payload)
        self.assertEqual(err, "A data é obrigatória.")

    def test_montar_agrupa_por_carro(self):
        p1 = self._pedido(91001, "REF-A")
        p2 = self._pedido(91002, "REF-B")
        self._tentativa(p1, carro=3)
        self._tentativa(p2, carro=3)
        payload, err = montar_relatorio_rotas(
            self.filial,
            {"data_tentativa": self.d0.isoformat(), "agrupamento": "carro"},
        )
        self.assertIsNone(err)
        self.assertEqual(payload["agrupamento"], "carro")
        self.assertEqual(len(payload["grupos"]), 1)
        self.assertEqual(payload["grupos"][0]["carro"], "3")
        self.assertEqual(payload["grupos"][0]["total"], 2)

    def test_gerar_xlsx_retorna_arquivo_zip(self):
        if importlib.util.find_spec("openpyxl") is None:
            self.skipTest("openpyxl não instalado")
        from openpyxl import load_workbook

        p = self._pedido(91003, "REF-X")
        self._tentativa(p, carro=7)
        payload, err = montar_relatorio_rotas(
            self.filial,
            {"data_tentativa": self.d0.isoformat()},
        )
        self.assertIsNone(err)
        data = gerar_xlsx_relatorio_rotas(payload)
        self.assertTrue(data.startswith(b"PK"))
        self.assertGreater(len(data), 100)

        wb = load_workbook(BytesIO(data))
        ws = wb.active
        # título + linha vazia + cabeçalho + 1 dado
        self.assertEqual(ws.cell(row=4, column=1).value, "Carro 7")
        self.assertEqual(ws.cell(row=4, column=3).value, "REF-X")

    def test_tem_incidencia_peso_pendente(self):
        p_pendente = self._pedido(91010, "REF-PESO")
        p_resolvido = self._pedido(91011, "REF-OK")
        p_outro_tipo = self._pedido(91012, "REF-OUTRO")
        self._tentativa(p_pendente, carro=1)
        self._tentativa(p_resolvido, carro=1)
        self._tentativa(p_outro_tipo, carro=1)

        Incidencia.objects.create(
            pedido=p_pendente,
            data=self.d0,
            origem="Cliente",
            tipo="Peso/Volume",
            resolvido=False,
        )
        Incidencia.objects.create(
            pedido=p_resolvido,
            data=self.d0,
            origem="Cliente",
            tipo="Peso/Volume",
            resolvido=True,
        )
        Incidencia.objects.create(
            pedido=p_outro_tipo,
            data=self.d0,
            origem="Cliente",
            tipo="Data/Horário",
            resolvido=False,
        )

        payload, err = montar_relatorio_rotas(
            self.filial,
            {"data_tentativa": self.d0.isoformat(), "agrupamento": "carro"},
        )
        self.assertIsNone(err)
        por_ref = {ln["pedido"]: ln for ln in payload["grupos"][0]["linhas"]}
        self.assertTrue(por_ref["REF-PESO"]["tem_incidencia_peso_pendente"])
        self.assertFalse(por_ref["REF-OK"]["tem_incidencia_peso_pendente"])
        self.assertFalse(por_ref["REF-OUTRO"]["tem_incidencia_peso_pendente"])

    def test_incidencia_com_motivo_bloqueante_nao_segue(self):
        p = self._pedido(91020, "REF-MOT")
        TentativaEntrega.objects.create(
            pedido=p,
            data_tentativa=self.d0,
            carro=1,
            periodo="MANHA",
            estado="Incidência",
            motivo_incidencia="Fora da zona",
        )
        payload, err = montar_relatorio_rotas(
            self.filial,
            {"data_tentativa": self.d0.isoformat(), "agrupamento": "carro"},
        )
        self.assertIsNone(err)
        linha = payload["grupos"][0]["linhas"][0]
        self.assertFalse(linha["segue_para_entrega"])
        self.assertTrue(linha["nao_segue_para_entrega"])
