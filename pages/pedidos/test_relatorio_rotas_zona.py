"""Relatório de Rotas por Zona — montagem de grupos e exportação XLSX."""

import importlib.util
from datetime import date, timedelta
from io import BytesIO

from django.test import TestCase
from django.utils import timezone

from pages.core.models import Pais
from pages.filial.models import Filial
from pages.pedidos.models import Incidencia, Pedido, TentativaEntrega
from pages.pedidos.services.relatorio_rotas_zona import (
    PERIODO_MAXIMO_ROTAS_ZONA_DIAS,
    SEM_ZONA_FILTRO,
    SEM_ZONA_LABEL,
    gerar_xlsx_relatorio_rotas_zona,
    listar_zonas_choices_filial,
    montar_relatorio_rotas_zona,
)
from pages.zona_entrega.models import ZonaEntrega, ZonaEntregaFaixaPostal


class RelatorioRotasZonaServiceTests(TestCase):
    d0 = date(2026, 7, 1)
    d1 = date(2026, 7, 10)

    @classmethod
    def setUpTestData(cls):
        cls.pais = Pais.objects.create(nome="PT-RRZ", sigla="PRT", codigo_tel="+351")
        cls.filial = Filial.objects.create(
            codigo="RRZN",
            nome="FIL RR ZONA",
            pais_atuacao=cls.pais,
            is_matriz=True,
        )
        cls.outra_filial = Filial.objects.create(
            codigo="RRZO",
            nome="FIL RR ZONA OUTRA",
            pais_atuacao=cls.pais,
            is_matriz=False,
        )
        cls.zona_alta = ZonaEntrega.objects.create(
            filial=cls.filial,
            codigo="ZNALTA",
            descricao="ZONA NORTE",
            prioridade=10,
            ativa=True,
        )
        cls.zona_baixa = ZonaEntrega.objects.create(
            filial=cls.filial,
            codigo="ZNBAIX",
            descricao="ZONA SUL",
            prioridade=1,
            ativa=True,
        )
        ZonaEntregaFaixaPostal.objects.create(
            zona_entrega=cls.zona_alta,
            tipo_intervalo="CP7",
            codigo_postal_inicial="4000-000",
            codigo_postal_final="4999-999",
            ativa=True,
            cp4_inicial="0000",
            cp4_final="0000",
            cp7_inicial_num=0,
            cp7_final_num=0,
        )
        ZonaEntregaFaixaPostal.objects.create(
            zona_entrega=cls.zona_baixa,
            tipo_intervalo="CP7",
            codigo_postal_inicial="1000-000",
            codigo_postal_final="1999-999",
            ativa=True,
            cp4_inicial="0000",
            cp4_final="0000",
            cp7_inicial_num=0,
            cp7_final_num=0,
        )

    def _pedido(self, id_vonzu: int, *, filial=None, pedido_ref: str = "", codpost: str = "", peso=None) -> Pedido:
        now = timezone.now()
        return Pedido.objects.create(
            filial=filial or self.filial,
            id_vonzu=id_vonzu,
            pedido=pedido_ref or str(id_vonzu),
            tipo="ENTREGA",
            nome_dest="Dest Teste",
            codpost_dest=codpost,
            peso=peso,
            criado=now,
            atualizacao=now,
        )

    def _tentativa(self, pedido: Pedido, dt=None, carro: int = 1) -> TentativaEntrega:
        return TentativaEntrega.objects.create(
            pedido=pedido,
            data_tentativa=dt or self.d0,
            carro=carro,
            periodo="MANHA",
            estado="created",
        )

    def _filtros(self, **extra):
        base = {
            "data_inicial": self.d0.isoformat(),
            "data_final": self.d1.isoformat(),
        }
        base.update(extra)
        return base

    def test_montar_exige_filial(self):
        payload, err = montar_relatorio_rotas_zona(None, self._filtros())
        self.assertIsNone(payload)
        self.assertIn("Filial ativa", err)

    def test_montar_exige_datas(self):
        payload, err = montar_relatorio_rotas_zona(self.filial, {})
        self.assertIsNone(payload)
        self.assertEqual(err, "A data inicial é obrigatória.")

        payload, err = montar_relatorio_rotas_zona(
            self.filial, {"data_inicial": self.d0.isoformat()}
        )
        self.assertIsNone(payload)
        self.assertEqual(err, "A data final é obrigatória.")

    def test_periodo_invertido(self):
        payload, err = montar_relatorio_rotas_zona(
            self.filial,
            {"data_inicial": self.d1.isoformat(), "data_final": self.d0.isoformat()},
        )
        self.assertIsNone(payload)
        self.assertIn("maior", err)

    def test_periodo_maximo_365(self):
        ini = date(2026, 1, 1)
        fim = ini + timedelta(days=PERIODO_MAXIMO_ROTAS_ZONA_DIAS + 1)
        payload, err = montar_relatorio_rotas_zona(
            self.filial,
            {"data_inicial": ini.isoformat(), "data_final": fim.isoformat()},
        )
        self.assertIsNone(payload)
        self.assertIn("365", err)

        fim_ok = ini + timedelta(days=PERIODO_MAXIMO_ROTAS_ZONA_DIAS)
        payload, err = montar_relatorio_rotas_zona(
            self.filial,
            {"data_inicial": ini.isoformat(), "data_final": fim_ok.isoformat()},
        )
        self.assertIsNone(err)
        self.assertIsNotNone(payload)

    def test_agrupa_por_zona_e_sem_zona_no_fim(self):
        p_norte = self._pedido(92001, pedido_ref="REF-N", codpost="4100-100")
        p_sul = self._pedido(92002, pedido_ref="REF-S", codpost="1200-100")
        p_sem = self._pedido(92003, pedido_ref="REF-X", codpost="")
        self._tentativa(p_norte)
        self._tentativa(p_sul)
        self._tentativa(p_sem)

        payload, err = montar_relatorio_rotas_zona(self.filial, self._filtros())
        self.assertIsNone(err)
        nomes = [g["zona"] for g in payload["grupos"]]
        self.assertEqual(nomes[0], "ZONA NORTE")
        self.assertEqual(nomes[1], "ZONA SUL")
        self.assertEqual(nomes[-1], SEM_ZONA_LABEL)
        self.assertEqual(payload["grupos"][0]["zona_id"], self.zona_alta.id)
        self.assertEqual(payload["grupos"][0]["total"], 1)

    def test_filtro_zona_exclui_outras_e_sem_zona(self):
        p_norte = self._pedido(92011, pedido_ref="REF-N2", codpost="4100-100")
        p_sul = self._pedido(92012, pedido_ref="REF-S2", codpost="1200-100")
        p_sem = self._pedido(92013, pedido_ref="REF-X2", codpost="")
        self._tentativa(p_norte)
        self._tentativa(p_sul)
        self._tentativa(p_sem)

        payload, err = montar_relatorio_rotas_zona(
            self.filial,
            self._filtros(zonas=[self.zona_alta.id]),
        )
        self.assertIsNone(err)
        self.assertEqual(len(payload["grupos"]), 1)
        self.assertEqual(payload["grupos"][0]["zona_id"], self.zona_alta.id)
        self.assertEqual(payload["grupos"][0]["linhas"][0]["pedido"], "REF-N2")

    def test_filtro_sem_zona_so_nao_resolvidas(self):
        p_norte = self._pedido(92014, pedido_ref="REF-N3", codpost="4100-100")
        p_sem = self._pedido(92015, pedido_ref="REF-X3", codpost="")
        self._tentativa(p_norte)
        self._tentativa(p_sem)

        payload, err = montar_relatorio_rotas_zona(
            self.filial,
            self._filtros(zonas=[SEM_ZONA_FILTRO]),
        )
        self.assertIsNone(err)
        self.assertEqual(len(payload["grupos"]), 1)
        self.assertIsNone(payload["grupos"][0]["zona_id"])
        self.assertEqual(payload["grupos"][0]["linhas"][0]["pedido"], "REF-X3")

    def test_filtro_zona_e_sem_zona_juntos(self):
        p_norte = self._pedido(92016, pedido_ref="REF-N4", codpost="4100-100")
        p_sul = self._pedido(92017, pedido_ref="REF-S4", codpost="1200-100")
        p_sem = self._pedido(92018, pedido_ref="REF-X4", codpost="")
        self._tentativa(p_norte)
        self._tentativa(p_sul)
        self._tentativa(p_sem)

        payload, err = montar_relatorio_rotas_zona(
            self.filial,
            self._filtros(zonas=[self.zona_alta.id, SEM_ZONA_FILTRO]),
        )
        self.assertIsNone(err)
        nomes = {g["zona"] for g in payload["grupos"]}
        self.assertEqual(nomes, {"ZONA NORTE", SEM_ZONA_LABEL})

    def test_tenant_outra_filial_nao_entra(self):
        p_ok = self._pedido(92021, pedido_ref="REF-OK", codpost="4100-100")
        p_outra = self._pedido(
            92022, filial=self.outra_filial, pedido_ref="REF-OUT", codpost="4100-100"
        )
        self._tentativa(p_ok)
        self._tentativa(p_outra)

        payload, err = montar_relatorio_rotas_zona(self.filial, self._filtros())
        self.assertIsNone(err)
        refs = [ln["pedido"] for g in payload["grupos"] for ln in g["linhas"]]
        self.assertIn("REF-OK", refs)
        self.assertNotIn("REF-OUT", refs)

    def test_fora_do_periodo_nao_entra(self):
        p = self._pedido(92031, pedido_ref="REF-FORA", codpost="4100-100")
        self._tentativa(p, dt=date(2026, 8, 1))
        payload, err = montar_relatorio_rotas_zona(self.filial, self._filtros())
        self.assertIsNone(err)
        self.assertEqual(payload["grupos"], [])

    def test_tentativa_posterior_marca_nao_segue(self):
        p = self._pedido(92041, pedido_ref="REF-POST", codpost="4100-100")
        self._tentativa(p, dt=self.d0)
        TentativaEntrega.objects.create(
            pedido=p,
            data_tentativa=self.d1,
            carro=1,
            periodo="TARDE",
            estado="created",
        )
        payload, err = montar_relatorio_rotas_zona(self.filial, self._filtros())
        self.assertIsNone(err)
        linhas = payload["grupos"][0]["linhas"]
        por_data = {ln["data_tentativa"]: ln for ln in linhas}
        self.assertTrue(por_data["01/07/2026"]["nao_segue_para_entrega"])
        self.assertFalse(por_data["10/07/2026"]["nao_segue_para_entrega"])

    def test_incidencia_peso_pendente(self):
        p = self._pedido(92051, pedido_ref="REF-PESO", codpost="4100-100")
        self._tentativa(p)
        Incidencia.objects.create(
            pedido=p,
            data=self.d0,
            origem="Cliente",
            tipo="Peso/Volume",
            resolvido=False,
        )
        payload, err = montar_relatorio_rotas_zona(self.filial, self._filtros())
        self.assertIsNone(err)
        self.assertTrue(payload["grupos"][0]["linhas"][0]["tem_incidencia_peso_pendente"])

    def test_agrupa_por_data_e_conta_carros_distintos(self):
        p1 = self._pedido(92071, pedido_ref="REF-C1", codpost="4100-100", peso=10)
        p2 = self._pedido(92072, pedido_ref="REF-C3", codpost="4100-100", peso=5)
        p3 = self._pedido(92073, pedido_ref="REF-C5", codpost="4100-100")
        p4 = self._pedido(92074, pedido_ref="REF-C1B", codpost="4100-100")
        p_dia2 = self._pedido(92075, pedido_ref="REF-D2", codpost="4100-100")
        self._tentativa(p1, dt=self.d0, carro=1)
        self._tentativa(p2, dt=self.d0, carro=3)
        self._tentativa(p3, dt=self.d0, carro=5)
        self._tentativa(p4, dt=self.d0, carro=1)
        self._tentativa(p_dia2, dt=self.d1, carro=2)

        payload, err = montar_relatorio_rotas_zona(self.filial, self._filtros())
        self.assertIsNone(err)
        dias = payload["grupos"][0]["dias"]
        self.assertEqual(len(dias), 2)
        self.assertEqual(dias[0]["data"], "01/07/2026")
        self.assertEqual(dias[0]["carros"], ["1", "3", "5"])
        self.assertEqual(dias[0]["total_carros"], 3)
        self.assertEqual(dias[0]["total"], 4)
        self.assertEqual(dias[1]["carros"], ["2"])
        self.assertEqual(dias[1]["total_carros"], 1)

    def test_listar_zonas_choices(self):
        choices = listar_zonas_choices_filial(self.filial)
        ids = [c["value"] for c in choices]
        self.assertEqual(ids[0], self.zona_alta.id)
        self.assertIn(self.zona_baixa.id, ids)
        self.assertEqual(listar_zonas_choices_filial(None), [])

    def test_gerar_xlsx_retorna_arquivo_zip(self):
        if importlib.util.find_spec("openpyxl") is None:
            self.skipTest("openpyxl não instalado")
        from openpyxl import load_workbook

        p = self._pedido(92061, pedido_ref="REF-XLS", codpost="4100-100")
        self._tentativa(p, carro=7)
        payload, err = montar_relatorio_rotas_zona(self.filial, self._filtros())
        self.assertIsNone(err)
        data = gerar_xlsx_relatorio_rotas_zona(payload)
        self.assertTrue(data.startswith(b"PK"))

        wb = load_workbook(BytesIO(data))
        ws_resumo = wb["Resumo"]
        self.assertEqual(ws_resumo.cell(row=4, column=1).value, "ZONA NORTE")
        self.assertEqual(ws_resumo.cell(row=4, column=2).value, "01/07/2026")
        self.assertEqual(ws_resumo.cell(row=4, column=3).value, 1)
        self.assertEqual(ws_resumo.cell(row=4, column=6).value, "7")
        self.assertEqual(ws_resumo.cell(row=4, column=7).value, 1)

        ws = wb["Detalhe"]
        self.assertIn("Carros 7", ws.cell(row=4, column=1).value)
        self.assertEqual(ws.cell(row=5, column=5).value, "REF-XLS")
        self.assertEqual(ws.cell(row=5, column=3).value, "7")
